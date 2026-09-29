from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
DOCX_PATH = OUT_DIR / "جدول_الاختبارات_المدرسية.docx"
LOGO_PATH = ROOT / "public" / "school-tasks-logo.jpeg"

FONT = "Arial"
BLUE = "1F5FBF"
BLUE_DARK = "173A7A"
BLUE_LIGHT = "EDF5FF"
PINK_LIGHT = "FCE7F3"
PINK_SOFT = "FDF2F8"
BORDER = "9DB7D9"
ROW_ALT = "F8FBFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=95, bottom=70, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))


def set_table_bidi(table):
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_row_height(row, height_cm, exact=False):
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = tr_pr.find(qn("w:trHeight"))
    if tr_height is None:
        tr_height = OxmlElement("w:trHeight")
        tr_pr.append(tr_height)
    tr_height.set(qn("w:val"), str(int(height_cm / 2.54 * 1440)))
    tr_height.set(qn("w:hRule"), "exact" if exact else "atLeast")


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def add_run(paragraph, text, size=10, bold=False, color="111827"):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")
    sz_cs = r_pr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        r_pr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    return run


def cell_text(cell, text="", size=10, bold=False, color="111827", align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph_rtl(paragraph, align)
    add_run(paragraph, text, size=size, bold=bold, color=color)


def add_spacer(doc, pts):
    paragraph = doc.add_paragraph()
    set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(pts)


def main():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.75)
    section.bottom_margin = Cm(0.55)
    section.left_margin = Cm(0.85)
    section.right_margin = Cm(0.85)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(10)

    header = doc.add_table(rows=1, cols=3)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    set_table_bidi(header)
    widths = [6.5, 4.2, 8.0]
    for cell, width in zip(header.rows[0].cells, widths):
        set_cell_width(cell, width)
        set_cell_border(cell, "FFFFFF", "0")
        set_cell_margins(cell, 20, 40, 20, 40)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    right = header.rows[0].cells[0]
    right.paragraphs[0].clear()
    for line, bold in [
        ("الجمهورية اليمنية", True),
        ("وزارة التربية والتعليم والبحث العلمي", False),
        ("مكتب التربية بأمانة العاصمة", False),
        ("اسم المدرسة: __________________", True),
    ]:
        p = right.add_paragraph() if right.paragraphs[0].text else right.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        add_run(p, line, 8.6, bold, BLUE_DARK)

    logo_cell = header.rows[0].cells[1]
    p = logo_cell.paragraphs[0]
    p.clear()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    if LOGO_PATH.exists():
        run = p.add_run()
        run.add_picture(str(LOGO_PATH), width=Cm(2.85))

    left = header.rows[0].cells[2]
    left.paragraphs[0].clear()
    for line in ("العام الدراسي: __________", "الفصل الدراسي: __________"):
        p = left.add_paragraph() if left.paragraphs[0].text else left.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        add_run(p, line, 9.2, True, BLUE_DARK)

    add_spacer(doc, 4)

    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.autofit = False
    set_table_bidi(title_table)
    title_cell = title_table.cell(0, 0)
    set_cell_width(title_cell, 19.0)
    set_cell_border(title_cell, BLUE, "8")
    set_cell_shading(title_cell, BLUE)
    set_row_height(title_table.rows[0], 0.78, exact=True)
    cell_text(title_cell, "جدول الاختبارات المدرسية", size=16, bold=True, color="FFFFFF")

    add_spacer(doc, 5)

    fields = doc.add_table(rows=1, cols=3)
    fields.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields.autofit = False
    set_table_bidi(fields)
    field_widths = [9.9, 5.1, 4.0]
    field_texts = [
        "اسم الجدول (الصف أو المادة): ______________________________",
        "الفترة / الشهر: __________________",
        "الشعبة: __________",
    ]
    for cell, width, text in zip(fields.rows[0].cells, field_widths, field_texts):
        set_cell_width(cell, width)
        set_cell_shading(cell, BLUE_LIGHT)
        set_row_height(fields.rows[0], 0.78, exact=True)
        cell_text(cell, text, size=8.5, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.RIGHT)

    add_spacer(doc, 6)

    schedule = doc.add_table(rows=15, cols=5)
    schedule.alignment = WD_TABLE_ALIGNMENT.CENTER
    schedule.autofit = False
    set_table_bidi(schedule)
    col_widths = [1.15, 2.65, 3.05, 5.15, 7.0]
    headers = ["م", "اليوم", "التاريخ", "اسم المادة", "المقرر"]
    for index, cell in enumerate(schedule.rows[0].cells):
        set_cell_width(cell, col_widths[index])
        set_cell_shading(cell, PINK_LIGHT if index == 0 else BLUE)
        cell_text(cell, headers[index], size=9.5, bold=True, color=BLUE_DARK if index == 0 else "FFFFFF")
    set_row_height(schedule.rows[0], 0.75, exact=True)

    for row_idx in range(1, 15):
        row = schedule.rows[row_idx]
        set_row_height(row, 0.76, exact=True)
        for col_idx, cell in enumerate(row.cells):
            set_cell_width(cell, col_widths[col_idx])
            fill = PINK_SOFT if col_idx == 0 else (ROW_ALT if row_idx % 2 == 0 else "FFFFFF")
            set_cell_shading(cell, fill)
            cell_text(cell, str(row_idx) if col_idx == 0 else "", size=9.5, bold=col_idx == 0, color=BLUE_DARK if col_idx == 0 else "111827")

    add_spacer(doc, 8)

    signatures = doc.add_table(rows=2, cols=3)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.autofit = False
    set_table_bidi(signatures)
    sig_titles = ["إعداد الجدول", "مسؤول الاختبارات", "مدير المدرسة"]
    for row in signatures.rows:
        set_row_height(row, 0.62, exact=True)
    for col, title in enumerate(sig_titles):
        top = signatures.cell(0, col)
        bottom = signatures.cell(1, col)
        set_cell_width(top, 6.33)
        set_cell_width(bottom, 6.33)
        set_cell_shading(top, BLUE_LIGHT)
        set_cell_shading(bottom, "FFFFFF")
        cell_text(top, title, size=9.3, bold=True, color=BLUE_DARK)
        cell_text(bottom, "الاسم والتوقيع: __________________", size=8.4, bold=False, color="374151")

    add_spacer(doc, 5)

    footer = doc.add_paragraph()
    set_paragraph_rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(footer, "وحدة الحاسوب / محمد فراس", size=8.2, bold=True, color="6B7280")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
